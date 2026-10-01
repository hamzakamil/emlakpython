from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from users.models import User, UserRole


class ExcelDatabaseApiTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="excel-admin", password="x", email="excel-admin@example.com"
        )
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def test_template_contains_mapping_and_selected_table(self):
        response = self.client.post(
            "/api/v1/database/excel/template/",
            {"tables": ["tenants_tenant"]},
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
        self.assertIn("Talimatlar", workbook.sheetnames)
        self.assertIn("__TabloEsleme", workbook.sheetnames)
        self.assertIn("tenants_tenant", workbook.sheetnames)

    def test_dry_run_does_not_write(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "tenants_tenant"
        sheet.append(["name", "slug", "is_active", "max_users", "max_projects", "max_storage_gb"])
        sheet.append(["Dry Run Firma", "dry-run-firma", True, 5, 1, 5])
        stream = BytesIO()
        workbook.save(stream)
        upload = SimpleUploadedFile("dry-run.xlsx", stream.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response = self.client.post(
            "/api/v1/database/excel/import/",
            {"workbook": upload, "dry_run": "true"},
            format="multipart",
        )
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["dry_run"])
        from tenants.models import Tenant
        self.assertFalse(Tenant.objects.filter(slug="dry-run-firma").exists())

    def test_export_returns_populated_workbook(self):
        response = self.client.post(
            "/api/v1/database/excel/export/",
            {"tables": ["tenants_tenant"], "fields": {"tenants_tenant": ["name", "slug"]}},
            format="json",
        )
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)), read_only=True)
        self.assertEqual([cell.value for cell in workbook["tenants_tenant"][1]], ["name", "slug"])
